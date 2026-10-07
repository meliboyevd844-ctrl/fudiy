import datetime
from pathlib import Path
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from config import REPORTS_DIR

WEEKDAY_UZ = {
    0: "Dushanba",
    1: "Seshanba",
    2: "Chorshanba",
    3: "Payshanba",
    4: "Juma",
    5: "Shanba",
    6: "Yakshanba"
}

def format_minutes_to_hours(minutes: int) -> str:
    """Format minutes to readable 'X soat Y daq' format."""
    if not minutes or minutes <= 0:
        return "-"
    hours = minutes // 60
    mins = minutes % 60
    if hours > 0 and mins > 0:
        return f"{hours} soat {mins} daq"
    elif hours > 0:
        return f"{hours} soat"
    else:
        return f"{mins} daq"


def generate_monthly_excel_report(records: List[Dict[str, Any]], year: int, month: int) -> Path:
    """
    Generate a professional Excel (.xlsx) file for the monthly attendance report.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    month_str = f"{year}-{month:02d}"
    ws.title = f"Hisobot {month_str}"

    # Styles
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="475569")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    
    status_keldi_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # light green
    status_kechikdi_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # light yellow
    status_kelmadi_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # light red
    row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    # Title Block
    ws.merge_cells("A1:H1")
    ws["A1"] = f"DO'KON XODIMLARINING OYLIK KELDI-KETDI HISOBOTI ({month_str})"
    ws["A1"].font = title_font
    ws["A1"].alignment = align_center

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Hisobot yaratilgan sana va vaqt: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = align_center

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[4].height = 26

    # Table Headers
    headers = [
        "№",
        "Xodim F.I.SH",
        "Telefon raqam",
        "Sana",
        "Hafta kuni",
        "Kelgan vaqti",
        "Ketgan vaqti",
        "Holati",
        "Ishlagan vaqti"
    ]

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    # Populate Data
    current_row = 5
    for idx, rec in enumerate(records, 1):
        # Calculate weekday
        rec_date_str = rec.get("date") or "-"
        weekday_name = "-"
        if rec_date_str != "-":
            try:
                d_obj = datetime.datetime.strptime(rec_date_str, "%Y-%m-%d")
                weekday_name = WEEKDAY_UZ.get(d_obj.weekday(), "-")
            except Exception:
                pass

        status_raw = rec.get("status") or "kelmadi"
        status_display = "Kelmadi"
        fill_color = status_kelmadi_fill
        if status_raw == "keldi":
            status_display = "Vaqtida keldi"
            fill_color = status_keldi_fill
        elif status_raw == "kechikdi":
            status_display = "Kechikdi"
            fill_color = status_kechikdi_fill
        elif rec.get("arrival_time") and not rec.get("departure_time"):
            status_display = "Ishlamoqda"
            fill_color = status_keldi_fill

        arrival_str = rec.get("arrival_time") or "-"
        departure_str = rec.get("departure_time") or "-"
        work_time_str = format_minutes_to_hours(rec.get("total_work_minutes", 0))

        row_data = [
            idx,
            rec.get("full_name", "-"),
            rec.get("phone", "-"),
            rec_date_str,
            weekday_name,
            arrival_str,
            departure_str,
            status_display,
            work_time_str
        ]

        ws.row_dimensions[current_row].height = 20
        for col_num, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.value = val
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

            if col_num in [1, 4, 5, 6, 7, 9]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

            # Status column color
            if col_num == 8:
                cell.fill = fill_color
                cell.font = Font(name="Calibri", size=10, bold=True)
            elif current_row % 2 == 0:
                cell.fill = row_alt_fill

        current_row += 1

    # Auto adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws.column_dimensions["B"].width = 28 # Full name
    ws.column_dimensions["C"].width = 18 # Phone

    # Save to file
    filename = f"Hisobot_{month_str}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    file_path = REPORTS_DIR / filename
    wb.save(file_path)
    return file_path
